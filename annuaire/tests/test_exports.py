import io

import openpyxl
import pytest

from annuaire.exports import EXPORT_COLUMNS, build_export_rows, build_persons_workbook
from annuaire.models import Person


@pytest.mark.django_db
def test_build_export_rows_orders_by_last_first_pk(person, other_person):
    rows = build_export_rows([person.pk, other_person.pk])
    assert rows == [
        ["Busson", "Alice", "alice@example.com", "", ""],
        ["Busson", "Bob", "bob@example.com", "", ""],
    ]


@pytest.mark.django_db
def test_build_export_rows_blanks_none_fields(accountless_person):
    rows = build_export_rows([accountless_person.pk])
    assert rows == [["Busson", "Charlie", "", "", ""]]
    assert "None" not in "".join(rows[0])


@pytest.mark.django_db
def test_build_export_rows_drops_unknown_ids(person):
    rows = build_export_rows([person.pk, 999999])
    assert len(rows) == 1


@pytest.mark.django_db
def test_build_export_rows_empty_ids_returns_empty():
    assert build_export_rows([]) == []


@pytest.mark.django_db
def test_build_export_rows_skips_explicitly_redacted_person(person, other_person):
    other_person.export_privacy = Person.ExportPrivacy.REDACT
    other_person.save()
    rows = build_export_rows([person.pk, other_person.pk])
    assert len(rows) == 1
    assert rows[0][0] == "Busson" and rows[0][1] == "Alice"


@pytest.mark.django_db
def test_build_export_rows_keeps_living_person_under_auto(person):
    # Living-by-default redaction does NOT apply to the Excel export -- only
    # the explicit "redact" opt-out does.
    rows = build_export_rows([person.pk])
    assert len(rows) == 1


def test_build_persons_workbook_header_and_rows():
    rows = [["Busson", "Alice", "alice@example.com", "0102030405", "1 rue des Fleurs"]]
    workbook_bytes = build_persons_workbook(rows)
    workbook = openpyxl.load_workbook(filename=io.BytesIO(workbook_bytes))
    sheet = workbook.active
    assert sheet.title == "Annuaire"
    assert [cell.value for cell in sheet[1]] == [header for header, _ in EXPORT_COLUMNS]
    assert [cell.value for cell in sheet[2]] == rows[0]


def test_build_persons_workbook_empty_rows_still_valid():
    workbook_bytes = build_persons_workbook([])
    workbook = openpyxl.load_workbook(filename=io.BytesIO(workbook_bytes))
    sheet = workbook.active
    assert sheet.max_row == 1
    assert [cell.value for cell in sheet[1]] == [header for header, _ in EXPORT_COLUMNS]


@pytest.mark.django_db
def test_export_columns_match_person_field_names(person):
    for _, attr in EXPORT_COLUMNS:
        assert hasattr(Person, attr)
