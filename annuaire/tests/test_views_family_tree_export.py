import io

import openpyxl
import pytest
from django.urls import reverse

LOGIN_URL = "/annuaire/login/"


@pytest.mark.django_db
def test_export_requires_login(client, person):
    response = client.get(reverse("genealogie-export"), {"ids": [person.pk]})
    assert response.status_code == 302
    assert LOGIN_URL in response["Location"]


@pytest.mark.django_db
def test_export_returns_xlsx_for_given_ids(auth_client, person, other_person):
    response = auth_client.get(reverse("genealogie-export"), {"ids": [person.pk, other_person.pk]})
    assert response.status_code == 200
    assert response["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert response["Content-Disposition"].startswith("attachment;")
    assert response["Content-Disposition"].endswith('.xlsx"')

    workbook = openpyxl.load_workbook(filename=io.BytesIO(response.content))
    sheet = workbook.active
    names = [row[1] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert set(names) == {"Alice", "Bob"}


@pytest.mark.django_db
def test_export_drops_unknown_and_non_integer_ids(auth_client, person):
    response = auth_client.get(reverse("genealogie-export"), {"ids": [person.pk, "not-an-id", 999999]})
    assert response.status_code == 200
    workbook = openpyxl.load_workbook(filename=io.BytesIO(response.content))
    sheet = workbook.active
    assert sheet.max_row == 2


@pytest.mark.django_db
def test_export_with_no_ids_returns_header_only(auth_client):
    response = auth_client.get(reverse("genealogie-export"))
    assert response.status_code == 200
    workbook = openpyxl.load_workbook(filename=io.BytesIO(response.content))
    assert workbook.active.max_row == 1
